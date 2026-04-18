from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

_BG = RGBColor(0x0F, 0x17, 0x2A)
_ACCENT = RGBColor(0x60, 0xA5, 0xFA)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_MUTED = RGBColor(0x94, 0xA3, 0xB8)
_EASY = RGBColor(0x34, 0xD3, 0x99)
_MEDIUM = RGBColor(0xFB, 0xBF, 0x24)
_HARD = RGBColor(0xF8, 0x71, 0x71)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


def _new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_slide(prs: Presentation):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _BG
    return slide


def _add_text(slide, text: str, left, top, width, height, size: int,
              bold=False, color=None, align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or _WHITE
    return txBox


def _add_paragraph(tf, text: str, size: int, color=None, bold=False,
                   indent_level=0, space_before=0):
    from pptx.util import Pt as _Pt
    p = tf.add_paragraph()
    p.level = indent_level
    p.space_before = Emu(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = _Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or _WHITE
    return p


def _accent_bar(slide, top):
    bar = slide.shapes.add_shape(1, Inches(0.5), top, Inches(1.2), Emu(18000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _ACCENT
    bar.line.fill.background()


def _slide_number(slide, num: int, total: int):
    _add_text(
        slide, f"{num} / {total}",
        Inches(11.8), Inches(7.1), Inches(1.3), Inches(0.3),
        size=9, color=_MUTED, align=PP_ALIGN.RIGHT,
    )


def generate_pptx(analysis: dict) -> bytes:
    slides_data = analysis.get("slides", [])
    prs = _new_prs()
    total = len(slides_data)

    for slide_info in slides_data:
        slide = _blank_slide(prs)
        num = slide_info.get("slideNumber", 1)
        title = slide_info.get("title", "")
        bullets = slide_info.get("bullets", [])
        note = slide_info.get("speakerNote", "")

        if num == 1:
            _build_title_slide(slide, analysis, note)
        else:
            _build_content_slide(slide, title, bullets, note, num, total)

        _slide_number(slide, num, total)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def generate_quick_pptx(analysis: dict) -> bytes:
    slides_data = analysis.get("quickSlides", []) or analysis.get("slides", [])[:5]
    prs = _new_prs()
    total = len(slides_data)

    for slide_info in slides_data:
        slide = _blank_slide(prs)
        num = slide_info.get("slideNumber", 1)
        title = slide_info.get("title", "")
        bullets = slide_info.get("bullets", [])
        note = slide_info.get("speakerNote", "")

        if num == 1:
            _build_title_slide(slide, analysis, note)
        else:
            _build_content_slide(slide, title, bullets, note, num, total)

        _slide_number(slide, num, total)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _build_title_slide(slide, analysis: dict, note: str):
    title = analysis.get("paperTitle", "") or analysis.get("title", "Research Paper")
    authors = analysis.get("authors", [])
    venue = analysis.get("venue", "")
    year = analysis.get("year", "")
    story = analysis.get("coreStory", {})

    _add_text(
        slide, title,
        Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.8),
        size=28, bold=True, color=_WHITE,
    )

    author_str = ", ".join(authors[:3]) + (" et al." if len(authors) > 3 else "")
    _add_text(
        slide, f"{author_str}  •  {venue}  •  {year}",
        Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.4),
        size=13, color=_MUTED,
    )

    bar = slide.shapes.add_shape(1, Inches(0.8), Inches(3.65), Inches(11.7), Emu(14000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x1E, 0x2D, 0x45)
    bar.line.fill.background()

    core = story.get("keyResult", story.get("soWhat", ""))
    if core:
        _add_text(
            slide, f'"{core}"',
            Inches(1.0), Inches(3.75), Inches(11.3), Inches(0.7),
            size=14, color=_ACCENT,
        )

    _add_speaker_note(slide, note)


def _build_content_slide(slide, title: str, bullets: list, note: str, num: int, total: int):
    _accent_bar(slide, Inches(1.05))

    _add_text(
        slide, title,
        Inches(1.9), Inches(0.75), Inches(10.9), Inches(0.65),
        size=22, bold=True, color=_WHITE,
    )

    if bullets:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.7), Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True

        first = True
        for bullet in bullets:
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            p.space_before = Emu(120000)
            run = p.add_run()
            run.text = f"›  {bullet}"
            run.font.size = Pt(16)
            run.font.color.rgb = _WHITE

    _add_speaker_note(slide, note)


def _add_speaker_note(slide, text: str):
    if not text:
        return
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = text
