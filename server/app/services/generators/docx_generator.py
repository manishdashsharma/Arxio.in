from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_heading_color(run, r: int, g: int, b: int):
    run.font.color.rgb = RGBColor(r, g, b)


def _add_horizontal_rule(doc: Document):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "3B82F6")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _section_heading(doc: Document, text: str):
    _add_horizontal_rule(doc)
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        run.font.size = Pt(14)
        run.font.bold = True


def _sub_heading(doc: Document, text: str):
    h = doc.add_heading(text, level=3)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x37, 0x41, 0xF5)
        run.font.size = Pt(12)


def _body_text(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)


def _bullet(doc: Document, text: str):
    p = doc.add_paragraph(text, style="List Bullet")
    p.style.font.size = Pt(11)


def _badge_paragraph(doc: Document, label: str, value: str, label_color: tuple):
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor(*label_color)
    label_run.font.size = Pt(11)
    val_run = p.add_run(value)
    val_run.font.size = Pt(11)


def generate_docx(analysis: dict) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title = doc.add_heading(analysis.get("paperTitle", "Research Paper"), level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        run.font.size = Pt(20)

    authors = ", ".join(analysis.get("authors", []))
    venue = analysis.get("venue", "")
    year = analysis.get("year", "")
    meta = doc.add_paragraph(f"{authors}  •  {venue}  •  {year}")
    meta.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    meta.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    story = analysis.get("coreStory", {})
    if story:
        _section_heading(doc, "Core Story")
        _badge_paragraph(doc, "Problem", story.get("problem", ""), (0xDC, 0x26, 0x26))
        _badge_paragraph(doc, "Solution", story.get("solution", ""), (0x05, 0x96, 0x69))
        _badge_paragraph(doc, "Key Result", story.get("keyResult", ""), (0x1D, 0x4E, 0xD8))
        _badge_paragraph(doc, "So What", story.get("soWhat", ""), (0x7C, 0x3A, 0xED))
        doc.add_paragraph()

    overview = analysis.get("overview", "")
    if overview:
        _section_heading(doc, "Overview")
        _body_text(doc, overview)
        doc.add_paragraph()

    problem = analysis.get("problemStatement", "")
    if problem:
        _section_heading(doc, "Problem Statement")
        _body_text(doc, problem)
        doc.add_paragraph()

    key_terms = analysis.get("keyTerms", [])
    if key_terms:
        _section_heading(doc, "Key Terms")
        for term in key_terms:
            _sub_heading(doc, term.get("term", ""))
            _body_text(doc, term.get("definition", ""))
            analogy = term.get("analogy", "")
            if analogy:
                p = doc.add_paragraph()
                run = p.add_run("Analogy: ")
                run.bold = True
                run.font.color.rgb = RGBColor(0x92, 0x4C, 0x0A)
                run.font.size = Pt(11)
                p.add_run(analogy).font.size = Pt(11)
        doc.add_paragraph()

    methodology = analysis.get("methodology", {})
    if methodology:
        _section_heading(doc, "Methodology")
        _body_text(doc, methodology.get("summary", ""))
        steps = methodology.get("steps", [])
        if steps:
            _sub_heading(doc, "Steps")
            for step in steps:
                _bullet(doc, step)
        why = methodology.get("whyThisApproach", "")
        if why:
            _sub_heading(doc, "Why This Approach")
            _body_text(doc, why)
        tools = methodology.get("tools", [])
        if tools:
            p = doc.add_paragraph()
            run = p.add_run("Tools & Frameworks: ")
            run.bold = True
            run.font.size = Pt(11)
            p.add_run(", ".join(tools)).font.size = Pt(11)
        datasets = methodology.get("datasets", [])
        if datasets:
            _sub_heading(doc, "Datasets")
            for d in datasets:
                _bullet(doc, d)
        doc.add_paragraph()

    results = analysis.get("results", {})
    if results:
        _section_heading(doc, "Results")
        _body_text(doc, results.get("summary", ""))
        findings = results.get("keyFindings", [])
        if findings:
            _sub_heading(doc, "Key Findings")
            for f in findings:
                _bullet(doc, f)
        comparison = results.get("comparisonToPriorWork", "")
        if comparison:
            _sub_heading(doc, "Comparison to Prior Work")
            _body_text(doc, comparison)
        meaning = results.get("whatTheNumbersMean", "")
        if meaning:
            _sub_heading(doc, "What the Numbers Mean")
            _body_text(doc, meaning)
        doc.add_paragraph()

    critical = analysis.get("criticalAnalysis", {})
    if critical:
        _section_heading(doc, "Critical Analysis")
        strengths = critical.get("strengths", [])
        if strengths:
            _sub_heading(doc, "Strengths")
            for s in strengths:
                _bullet(doc, s)
        weaknesses = critical.get("weaknesses", [])
        if weaknesses:
            _sub_heading(doc, "Weaknesses")
            for w in weaknesses:
                _bullet(doc, w)
        open_qs = critical.get("openQuestions", [])
        if open_qs:
            _sub_heading(doc, "Open Questions")
            for q in open_qs:
                _bullet(doc, q)
        future = critical.get("futureWork", "")
        if future:
            _sub_heading(doc, "Future Work")
            _body_text(doc, future)
        change = critical.get("whatYouWouldChange", "")
        if change:
            _sub_heading(doc, "What I Would Change")
            _body_text(doc, change)
        doc.add_paragraph()

    qa = analysis.get("qaPrep", [])
    if qa:
        _section_heading(doc, "Q&A Preparation")
        difficulty_colors = {
            "easy": (0x05, 0x96, 0x69),
            "medium": (0xD9, 0x77, 0x06),
            "hard": (0xDC, 0x26, 0x26),
        }
        for item in qa:
            difficulty = item.get("difficulty", "medium")
            color = difficulty_colors.get(difficulty, (0x37, 0x41, 0x51))
            _badge_paragraph(doc, difficulty.upper(), f"  {item.get('question', '')}", color)
            _body_text(doc, item.get("answer", ""))
            why = item.get("whyTheyAskThis", "")
            if why:
                p = doc.add_paragraph()
                run = p.add_run("Why they ask this: ")
                run.italic = True
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                run.font.size = Pt(10)
                p.add_run(why).font.size = Pt(10)
            doc.add_paragraph()

    cheat = analysis.get("cheatSheet", {})
    if cheat:
        _section_heading(doc, "Cheat Sheet")
        key_numbers = cheat.get("keyNumbers", [])
        if key_numbers:
            _sub_heading(doc, "Key Numbers")
            for n in key_numbers:
                _bullet(doc, n)
        talking = cheat.get("talkingPoints", [])
        if talking:
            _sub_heading(doc, "Talking Points")
            for t in talking:
                _bullet(doc, t)
        fallbacks = cheat.get("fallbackAnswers", [])
        if fallbacks:
            _sub_heading(doc, "Fallback Answers")
            for fb in fallbacks:
                p = doc.add_paragraph()
                run = p.add_run(fb.get("trigger", "") + ":  ")
                run.bold = True
                run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
                run.font.size = Pt(11)
                p.add_run(fb.get("say", "")).font.size = Pt(11)
        do_not = cheat.get("doNotSay", [])
        if do_not:
            _sub_heading(doc, "Do NOT Say")
            for d in do_not:
                p = doc.add_paragraph()
                run = p.add_run("✗  ")
                run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
                run.font.size = Pt(11)
                p.add_run(d).font.size = Pt(11)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
